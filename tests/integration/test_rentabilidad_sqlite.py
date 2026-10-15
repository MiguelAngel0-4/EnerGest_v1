"""
Pruebas de integración de la rentabilidad sobre SQLite real (Actividad 4.4).

Ruta: tests/integration/test_rentabilidad_sqlite.py

Escenario de más de un año, con un reloj que avanza:
- A (Perkins, $10.000.000): alquilada y luego con un mantenimiento costoso -> Evaluar reemplazo.
- B (Cummins, sin valor de compra): solo un gasto reciente -> Vigilar (con advertencia).
- C (Kohler): sin movimientos -> Sin datos suficientes.
- D (Generac, $5.000.000): alquilada (saldo sin cobrar) y luego vendida.
"""

from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.business.models.comercial import DatosContrato, DatosVenta, ModalidadAlquiler
from src.business.models.gasto import DatosGasto
from src.business.models.mantenimiento import DatosMantenimiento, LineaInsumo, TipoMantenimiento
from src.business.models.planta import DatosPlanta
from src.business.models.rentabilidad import NivelIndicador as N
from src.business.models.rentabilidad import Periodo
from src.business.services.comercial_service import ComercialService
from src.business.services.consecutivo_service import ConsecutivoService
from src.business.services.gasto_service import GastoService
from src.business.services.mantenimiento_service import MantenimientoService
from src.business.services.planta_service import PlantaService
from src.business.services.rentabilidad_service import RentabilidadService
from src.config import settings
from src.infrastructure.database.connection import DatabaseManager
from src.infrastructure.reports.excel_reporte import ExportadorExcel
from src.infrastructure.repositories.sqlite_unidad_de_trabajo import crear_fabrica_uow
from src.shared.paths import get_resource_path

DIA = ModalidadAlquiler.DIA


class Reloj:
    hoy = date(2025, 6, 1)


@pytest.fixture(scope="module")
def escenario(tmp_path_factory: pytest.TempPathFactory):
    db = DatabaseManager(tmp_path_factory.mktemp("rentabilidad") / "r.db")
    db.initialize_schema(get_resource_path(settings.SCHEMA_RELATIVE_PATH))
    reloj = Reloj()
    fabrica = crear_fabrica_uow(db)
    hoy = lambda: reloj.hoy  # noqa: E731
    plantas = PlantaService(fabrica, ConsecutivoService(), hoy)
    gastos = GastoService(fabrica, hoy)
    mantenimientos = MantenimientoService(fabrica, hoy)
    comercial = ComercialService(fabrica, plantas, hoy)
    cats = {c.nombre: c.id for c in gastos.listar_categorias()}
    cliente = comercial.registrar_cliente("Constructora del Pacífico").id

    def planta(marca: str, valor: int | None) -> int:
        datos = DatosPlanta(
            marca=marca, potencia_kva=60, horometro_inicial=1000, valor_compra=valor
        )
        return plantas.registrar(datos).id

    a, b, c, d = (
        planta("Perkins", 10_000_000),
        planta("Cummins", None),
        planta("Kohler", 1),
        planta("Generac", 5_000_000),
    )

    # A: alquiler de 10 días cobrado, un gasto, uno anulado y, un año después, un mantenimiento caro
    comercial.alquilar(a, DatosContrato(cliente, DIA, 100_000), horometro=1020)
    reloj.hoy = date(2025, 6, 11)
    comercial.devolver(a, horometro=1100, cobrar_saldo=True)
    reloj.hoy = date(2025, 7, 1)
    gastos.registrar_gasto(a, DatosGasto(cats["Repuestos"], reloj.hoy, "Alternador", 500_000))
    anulado = gastos.registrar_gasto(a, DatosGasto(cats["Otros"], reloj.hoy, "Duplicado", 999_999))
    gastos.anular_gasto(anulado.id, "Registrado dos veces")

    # D: alquiler con saldo sin cobrar y luego venta
    reloj.hoy = date(2026, 1, 5)
    comercial.alquilar(d, DatosContrato(cliente, DIA, 50_000), horometro=1010)
    reloj.hoy = date(2026, 1, 15)
    comercial.devolver(d, horometro=1060)  # Saldo de $ 500.000 sin cobrar
    reloj.hoy = date(2026, 2, 1)
    comercial.vender(d, DatosVenta(cliente, 3_000_000))

    reloj.hoy = date(2026, 9, 1)
    mantenimientos.registrar(
        a,
        DatosMantenimiento(reloj.hoy, TipoMantenimiento.CORRECTIVO, 1150, "Reparación del motor"),
        [LineaInsumo(cats["Repuestos"], "Kit de reparación", 7_000_000)],
    )
    # B: un gasto reciente
    reloj.hoy = date(2026, 10, 1)
    gastos.registrar_gasto(b, DatosGasto(cats["Mano de obra"], reloj.hoy, "Revisión", 200_000))

    reloj.hoy = date(2026, 10, 8)
    servicio = RentabilidadService(fabrica, hoy, umbral_desgaste=60, meses_perdida=12)
    return servicio, {"A": a, "B": b, "C": c, "D": d}


def test_balance_de_toda_la_vida(escenario) -> None:
    servicio, p = escenario
    balance = servicio.balance(p["A"])

    assert balance.totales.ingresos_alquiler == 1_000_000
    assert balance.totales.gastos == 7_500_000  # El gasto anulado no cuenta
    assert balance.totales.resultado_operacion == -6_500_000
    assert balance.recuperacion == -16_500_000  # -6.500.000 - 10.000.000 de compra
    assert balance.horas_uso_alquiler == 80
    assert balance.ingreso_por_hora == 12_500
    assert balance.gastos_por_categoria == (("Repuestos", 7_500_000),)


def test_el_periodo_filtra_los_movimientos(escenario) -> None:
    servicio, p = escenario
    balance = servicio.balance(p["A"], Periodo(date(2025, 1, 1), date(2025, 12, 31)))
    assert (balance.totales.ingresos_alquiler, balance.totales.gastos) == (1_000_000, 500_000)


def test_semaforo_usa_toda_la_vida_aunque_se_consulte_otro_periodo(escenario) -> None:
    servicio, p = escenario
    solo_2025 = servicio.balance(p["A"], Periodo(date(2025, 1, 1), date(2025, 12, 31)))
    assert solo_2025.indicador.nivel is N.EVALUAR_REEMPLAZO  # 75 % de desgaste y pérdida reciente
    assert solo_2025.indicador.desgaste_pct == pytest.approx(75)


def test_semaforo_de_cada_planta(escenario) -> None:
    servicio, p = escenario
    niveles = {f.planta.id: f.indicador.nivel for f in servicio.resumen().filas}
    assert niveles[p["A"]] is N.EVALUAR_REEMPLAZO
    assert niveles[p["B"]] is N.VIGILAR
    assert niveles[p["C"]] is N.SIN_DATOS
    indicador_b = servicio.balance(p["B"]).indicador
    assert any("Sin valor de compra" in r for r in indicador_b.razones)


def test_planta_vendida_recupera_con_su_venta(escenario) -> None:
    servicio, p = escenario
    balance = servicio.balance(p["D"])
    assert balance.totales.ingresos_venta == 3_000_000
    assert balance.saldo_por_cobrar == 500_000
    # Recuperación = 0 cobrado - 0 gastos + 3.000.000 venta - 5.000.000 compra
    assert balance.recuperacion == -2_000_000


def test_resumen_financiero(escenario) -> None:
    servicio, _ = escenario
    resumen = servicio.resumen()
    assert resumen.ingresos == 1_000_000 + 3_000_000
    assert resumen.gastos == 7_500_000 + 200_000
    assert resumen.resultado == -3_700_000
    assert resumen.saldo_por_cobrar == 500_000
    assert [m.mes for m in resumen.meses] == ["2025-06", "2025-07", "2026-02", "2026-09", "2026-10"]
    en_operacion = servicio.resumen(incluir_fuera_de_operacion=False)
    assert len(en_operacion.filas) == 3  # D está vendida


def test_ventana_reciente_del_semaforo(escenario) -> None:
    servicio, _ = escenario
    assert servicio.periodo_reciente() == Periodo(date(2025, 10, 9), date(2026, 10, 8))


def test_movimientos_con_el_formato_de_la_aplicacion(escenario) -> None:
    servicio, _ = escenario
    movimientos = servicio.movimientos()
    # Vigentes: cobro del alquiler de A, repuesto de A, venta de D, insumo del
    # mantenimiento de A y gasto de B. La devolución de D quedó sin cobrar (no hay
    # ingreso) y el gasto duplicado está anulado.
    assert len(movimientos) == 5
    assert movimientos[0].planta == "PE-001 · Perkins"
    assert {m.tipo for m in movimientos} == {"Ingreso", "Gasto"}
    assert all(m.descripcion != "Duplicado" for m in movimientos)  # Anulado excluido


def test_exportar_a_excel(escenario, tmp_path: Path) -> None:
    servicio, _ = escenario
    ruta = servicio.exportar(Periodo(), ExportadorExcel(), tmp_path / "reporte.xlsx")

    libro = load_workbook(ruta)
    assert libro.sheetnames == ["Resumen", "Por mes", "Movimientos"]
    resumen = libro["Resumen"]
    assert resumen["B5"].value == 4_000_000  # Ingresos: un número, no un texto
    assert resumen["B7"].value == -3_700_000
    assert "$" in resumen["B5"].number_format
    semaforos = [resumen.cell(fila, 10).value for fila in range(11, 15)]
    assert "Evaluar reemplazo" in semaforos
    assert libro["Movimientos"].max_row == 6  # Encabezado + 5 movimientos


def test_hoja_de_vida_v2_con_alquileres_y_resumen_financiero(escenario, tmp_path: Path) -> None:
    pypdf = pytest.importorskip("pypdf")
    from src.business.models.hoja_vida import DatosEmpresa
    from src.business.services.hoja_vida_service import HojaVidaService
    from src.infrastructure.reports.pdf_hoja_vida import GeneradorPdfHojaVida

    servicio, p = escenario
    hojas = HojaVidaService(
        servicio._uow_factory,
        GeneradorPdfHojaVida(),
        lambda: DatosEmpresa("WG ENERGIA"),
        "0.8.0",
        rentabilidad=servicio,
    )
    ruta = hojas.exportar(p["A"], tmp_path / "hoja.pdf")

    paginas = [pagina.extract_text() for pagina in pypdf.PdfReader(ruta).pages]
    texto = " ".join(" ".join(paginas).split())  # Une los renglones de las celdas
    assert "8. HISTORIAL DE ALQUILERES" in texto and "Constructora del Pacífico" in texto
    assert "9. RESUMEN FINANCIERO" in texto
    assert "Evaluar reemplazo" in texto and "75 %" in texto
    assert "Elaborado por" in paginas[-1] and "RESUMEN FINANCIERO" in paginas[-1]
    assert "01/06/2025" in texto  # Defecto corregido: la fecha ya no se parte en dos renglones
