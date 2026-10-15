"""
Pruebas de la pantalla Reportes financieros (Actividad 4.4).

Ruta: tests/ui/test_reportes_ui.py
"""

from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook
from PySide6.QtCore import QSettings, Qt
from PySide6.QtWidgets import QApplication, QLabel

from src.business.models.estado_planta import EstadoPlanta as E
from src.business.models.gasto import DatosGasto
from src.business.models.planta import DatosPlanta
from src.business.services.consecutivo_service import ConsecutivoService
from src.business.services.gasto_service import GastoService
from src.business.services.planta_service import PlantaService
from src.business.services.rentabilidad_service import RentabilidadService
from src.config import settings
from src.infrastructure.database.connection import DatabaseManager
from src.infrastructure.reports.excel_reporte import ExportadorExcel
from src.infrastructure.repositories.sqlite_unidad_de_trabajo import crear_fabrica_uow
from src.presentation.controllers.reportes_controller import ReportesController
from src.presentation.table_models.reporte_table_model import Col, FilaTablaReporte
from src.presentation.views.main_window import MainWindow
from src.presentation.views.reportes.reportes_view import ReportesView, Tarjeta
from src.shared.paths import get_resource_path


def _fila(planta_id: int, ingresos: int, nivel: str) -> FilaTablaReporte:
    return FilaTablaReporte(
        planta_id,
        f"PE-00{planta_id}",
        planta_id,
        "Planta",
        "Disponible",
        ingresos,
        0,
        ingresos,
        None,
        nivel,
        nivel,
        "",
    )


def _columna(vista: ReportesView, columna: int) -> list[str]:
    modelo = vista.tabla.model()
    return [modelo.index(f, columna).data() for f in range(modelo.rowCount())]


def test_ordena_por_valor_y_no_por_texto(qapp: QApplication) -> None:
    """ "$ 9.000" debe quedar ANTES que "$ 10.000" (como texto quedaría después)."""
    vista = ReportesView()
    vista.mostrar_filas([_fila(1, 10_000, "RENTABLE"), _fila(2, 9_000, "RENTABLE")])
    vista.tabla.sortByColumn(Col.INGRESOS, Qt.SortOrder.AscendingOrder)
    assert _columna(vista, Col.INGRESOS) == ["$ 9.000", "$ 10.000"]


def test_semaforo_ordena_de_lo_mas_grave_a_lo_menos(qapp: QApplication) -> None:
    vista = ReportesView()
    vista.mostrar_filas(
        [_fila(1, 0, "RENTABLE"), _fila(2, 0, "EVALUAR_REEMPLAZO"), _fila(3, 0, "VIGILAR")]
    )
    vista.tabla.sortByColumn(Col.SEMAFORO, Qt.SortOrder.AscendingOrder)
    assert _columna(vista, Col.SEMAFORO) == ["EVALUAR_REEMPLAZO", "VIGILAR", "RENTABLE"]


def test_tarjetas_y_grafico(qapp: QApplication) -> None:
    vista = ReportesView()
    vista.show()
    vista.mostrar_tarjetas([Tarjeta("Resultado", "-$ 5.000", negativo=True)])
    vista.mostrar_meses(["oct 2026"], [2_000_000], [500_000])
    QApplication.processEvents()
    valor = next(e for e in vista.findChildren(QLabel) if e.text() == "-$ 5.000")
    assert "color" in valor.styleSheet()
    series = vista._grafico.chart().series()
    assert len(series) == 1 and [s.label() for s in series[0].barSets()] == ["Ingresos", "Gastos"]

    vista.mostrar_meses([], [], [])  # Período sin movimientos
    assert vista._sin_movimientos.isVisible() and not vista._grafico.isVisible()


def test_el_menu_de_reportes_se_habilita_al_registrar_su_vista(qapp: QApplication) -> None:
    ventana = MainWindow("EnerGest", "0.8.0")
    assert not ventana._boton_reportes.isEnabled()
    ventana.agregar_vista("inventario", ReportesView())
    ventana.agregar_vista("reportes", ReportesView())
    assert ventana._boton_reportes.isEnabled()
    ventana._boton_reportes.click()
    assert ventana._contenedor.currentWidget() is ventana._vistas["reportes"]


# --- Controlador sobre SQLite real ------------------------------------------------


@pytest.fixture
def entorno(qapp: QApplication, tmp_path: Path):
    db = DatabaseManager(tmp_path / "reportes.db")
    db.initialize_schema(get_resource_path(settings.SCHEMA_RELATIVE_PATH))
    fabrica = crear_fabrica_uow(db)
    plantas = PlantaService(fabrica, ConsecutivoService())
    gastos = GastoService(fabrica)
    repuestos = next(c.id for c in gastos.listar_categorias() if c.nombre == "Repuestos")
    activa = plantas.registrar(DatosPlanta(marca="Perkins", potencia_kva=60)).id
    vendida = plantas.registrar(DatosPlanta(marca="Cummins", potencia_kva=50)).id
    gastos.registrar_gasto(activa, DatosGasto(repuestos, date.today(), "Alternador", 300_000))
    plantas.cambiar_estado(vendida, E.VENDIDA, motivo="Venta")
    abiertas: list[tuple[int, str]] = []
    vista = ReportesView()
    preferencias = QSettings(str(tmp_path / "prefs.ini"), QSettings.Format.IniFormat)
    controlador = ReportesController(
        RentabilidadService(fabrica),
        ExportadorExcel(),
        vista,
        abrir_ficha=lambda pid, pestana: abiertas.append((pid, pestana)),
        preferencias=preferencias,
    )
    return controlador, vista, activa, abiertas


def test_recargar_y_filtrar_plantas_fuera_de_operacion(entorno) -> None:
    controlador, vista, _activa, _ = entorno
    controlador.recargar()
    assert vista.tabla.model().rowCount() == 2
    assert "$ 300.000" in [e.text() for e in vista.findChildren(QLabel)]  # Tarjeta de gastos

    vista._incluir.setChecked(False)  # Dispara la recarga
    assert vista.tabla.model().rowCount() == 1


def test_exportar_lo_que_se_ve(entorno, tmp_path: Path, monkeypatch) -> None:
    controlador, vista, _activa, _ = entorno
    destino = tmp_path / "salida" / "reporte"  # Sin extensión: se agrega .xlsx
    monkeypatch.setattr(
        "PySide6.QtWidgets.QFileDialog.getSaveFileName", lambda *a, **k: (str(destino), "")
    )
    monkeypatch.setattr(vista, "preguntar_abrir", lambda ruta: False)
    vista._incluir.setChecked(False)

    controlador.exportar()

    libro = load_workbook(tmp_path / "salida" / "reporte.xlsx")
    filas = [libro["Resumen"].cell(f, 2).value for f in range(11, 13)]
    assert filas == ["Perkins", None]  # Solo la planta en operación, como en pantalla


def test_doble_clic_abre_el_balance_de_la_planta(entorno) -> None:
    controlador, vista, activa, abiertas = entorno
    controlador.recargar()
    fila = next(f for f in range(vista.modelo.rowCount()) if vista.modelo.planta_id(f) == activa)
    vista.tabla.doubleClicked.emit(vista.tabla.model().mapFromSource(vista.modelo.index(fila, 0)))
    assert abiertas == [(activa, "Balance")]
