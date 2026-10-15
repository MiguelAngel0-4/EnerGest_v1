"""
Exportador del reporte financiero a Excel (.xlsx) con openpyxl.

Ruta: src/infrastructure/reports/excel_reporte.py

Tres hojas:
1. Resumen: totales del período y una fila por planta, con su semáforo.
2. Por mes: ingresos, gastos y resultado de cada mes.
3. Movimientos: el detalle de cada ingreso y gasto vigente.

Los valores se escriben como NÚMEROS con formato de moneda (no como texto),
para que la empresa pueda sumar, filtrar y graficar en Excel.
"""

import logging
import os
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Any, Final

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src.business.exceptions import DocumentoError
from src.business.models.rentabilidad import Movimiento, Periodo, ResumenFinanciero
from src.shared.formatters import formatear_consecutivo

logger = logging.getLogger(__name__)

MONEDA: Final[str] = '"$" #,##0;[Red]-"$" #,##0'
FECHA: Final[str] = "dd/mm/yyyy"
PORCENTAJE: Final[str] = '0" %"'
_AZUL: Final[str] = "1F3A5F"
_ENCABEZADO = Font(bold=True, color="FFFFFF")
_RELLENO = PatternFill("solid", start_color=_AZUL, end_color=_AZUL)
_TITULO = Font(bold=True, size=14, color=_AZUL)


class ExportadorExcel:
    """Cumple el contrato IExportadorReporte."""

    def exportar(
        self, resumen: ResumenFinanciero, movimientos: list[Movimiento], destino: Path
    ) -> Path:
        temporal: Path | None = None
        try:
            libro = Workbook()
            self._hoja_resumen(libro.active, resumen)
            self._hoja_meses(libro.create_sheet("Por mes"), resumen)
            self._hoja_movimientos(libro.create_sheet("Movimientos"), movimientos)

            # Escritura segura, como el PDF: temporal y luego reemplazo atómico.
            destino.parent.mkdir(parents=True, exist_ok=True)
            descriptor, nombre = tempfile.mkstemp(suffix=".xlsx", dir=destino.parent)
            os.close(descriptor)
            temporal = Path(nombre)
            libro.save(temporal)
            os.replace(temporal, destino)
            return destino
        except OSError as exc:
            logger.exception("No se pudo escribir el reporte en %s", destino)
            raise DocumentoError(
                f"No se pudo guardar el archivo en:\n{destino}\n\n"
                "Si está abierto en Excel, ciérrelo e intente de nuevo."
            ) from exc
        finally:
            if temporal is not None and temporal.exists():
                temporal.unlink(missing_ok=True)

    # ------------------------------------------------------------------ #
    # Hojas
    # ------------------------------------------------------------------ #

    def _hoja_resumen(self, hoja: Worksheet, resumen: ResumenFinanciero) -> None:
        hoja.title = "Resumen"
        hoja["A1"] = "EnerGest – Reporte financiero"
        hoja["A1"].font = _TITULO
        hoja["A2"] = f"Período: {_texto_periodo(resumen.periodo)}"
        hoja["A3"] = f"Generado el {datetime.now():%d/%m/%Y %H:%M}"

        tarjetas = [
            ("Ingresos", resumen.ingresos),
            ("Gastos", resumen.gastos),
            ("Resultado", resumen.resultado),
            ("Saldo por cobrar (alquileres cerrados)", resumen.saldo_por_cobrar),
        ]
        for i, (nombre, valor) in enumerate(tarjetas, start=5):
            hoja.cell(i, 1, nombre).font = Font(bold=True)
            hoja.cell(i, 2, valor).number_format = MONEDA

        encabezados = [
            "N°",
            "Marca",
            "Modelo",
            "Estado",
            "Ingresos de operación",
            "Venta",
            "Gastos",
            "Resultado de operación",
            "Desgaste",
            "Semáforo",
            "Razones",
        ]
        fila_inicio = 10
        self._encabezado(hoja, fila_inicio, encabezados)
        for i, fila in enumerate(resumen.filas, start=fila_inicio + 1):
            planta, totales, indicador = fila.planta, fila.totales, fila.indicador
            valores: list[Any] = [
                formatear_consecutivo(planta.numero_consecutivo),
                planta.datos.marca,
                planta.datos.modelo or "",
                planta.estado.etiqueta,
                totales.ingresos_operacion,
                totales.ingresos_venta,
                totales.gastos,
                totales.resultado_operacion,
                round(indicador.desgaste_pct) if indicador.desgaste_pct is not None else None,
                indicador.nivel.etiqueta,
                " ".join(indicador.razones),
            ]
            for j, valor in enumerate(valores, start=1):
                celda = hoja.cell(i, j, valor)
                if 5 <= j <= 8:
                    celda.number_format = MONEDA
                elif j == 9:
                    celda.number_format = PORCENTAJE
        self._ajustar(
            hoja, fila_inicio, len(resumen.filas), [8, 14, 12, 16, 20, 14, 14, 22, 10, 18, 60]
        )

    def _hoja_meses(self, hoja: Worksheet, resumen: ResumenFinanciero) -> None:
        self._encabezado(hoja, 1, ["Mes", "Ingresos", "Gastos", "Resultado"])
        for i, mes in enumerate(resumen.meses, start=2):
            hoja.cell(i, 1, mes.mes)
            for j, valor in enumerate([mes.ingresos, mes.gastos, mes.ingresos - mes.gastos], 2):
                hoja.cell(i, j, valor).number_format = MONEDA
        self._ajustar(hoja, 1, len(resumen.meses), [10, 16, 16, 16])

    def _hoja_movimientos(self, hoja: Worksheet, movimientos: list[Movimiento]) -> None:
        self._encabezado(
            hoja, 1, ["Fecha", "Planta", "Tipo", "Concepto", "Descripción", "Documento", "Valor"]
        )
        for i, m in enumerate(movimientos, start=2):
            hoja.cell(i, 1, m.fecha).number_format = FECHA
            for j, valor in enumerate(
                [m.planta, m.tipo, m.concepto, m.descripcion, m.documento], 2
            ):
                hoja.cell(i, j, valor)
            hoja.cell(i, 7, m.valor).number_format = MONEDA
        self._ajustar(hoja, 1, len(movimientos), [12, 28, 10, 20, 44, 14, 14])

    # ------------------------------------------------------------------ #
    # Formato
    # ------------------------------------------------------------------ #

    @staticmethod
    def _encabezado(hoja: Worksheet, fila: int, textos: list[str]) -> None:
        for j, texto in enumerate(textos, start=1):
            celda = hoja.cell(fila, j, texto)
            celda.font = _ENCABEZADO
            celda.fill = _RELLENO
            celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    @staticmethod
    def _ajustar(hoja: Worksheet, fila_encabezado: int, filas: int, anchos: list[int]) -> None:
        """Anchos de columna, encabezado fijo al desplazarse y filtros automáticos."""
        for j, ancho in enumerate(anchos, start=1):
            hoja.column_dimensions[get_column_letter(j)].width = ancho
        hoja.freeze_panes = hoja.cell(fila_encabezado + 1, 1)
        ultima = get_column_letter(len(anchos))
        hoja.auto_filter.ref = f"A{fila_encabezado}:{ultima}{fila_encabezado + max(filas, 1)}"


def _texto_periodo(periodo: Periodo) -> str:
    if periodo.es_todo:
        return "Todo el historial"
    desde = periodo.desde.strftime("%d/%m/%Y") if periodo.desde else "el inicio"
    hasta = periodo.hasta.strftime("%d/%m/%Y") if periodo.hasta else "hoy"
    return f"Del {desde} al {hasta}"
